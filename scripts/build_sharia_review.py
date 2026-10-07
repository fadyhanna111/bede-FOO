"""Build a source-linked Murabaha capability review and test execution register.

Read only the preserved DOCX/XLSX sources. The outputs are review aids, not
implementation, Sharia approval, or completed test evidence.
"""

from collections import Counter
from pathlib import Path
import csv
import hashlib
import json
import re

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "documentation/sharia-murabaha/sources"
OUTPUT = ROOT / "planning/sharia-murabaha-2026-10-07"
OUTPUT.mkdir(parents=True, exist_ok=True)


def clean(value):
    return " ".join(str(value if value is not None else "").replace("\t", " ").split())


def write_tsv(path, rows, fields):
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows({key: clean(row.get(key, "")) for key in fields} for row in rows)


manifest = json.loads((SOURCE.parent / "source-manifest.json").read_text())
for entry in manifest["files"]:
    saved = ROOT / entry["project_path"]
    assert saved.is_file(), saved
    assert hashlib.sha256(saved.read_bytes()).hexdigest() == entry["sha256"], saved

assessment = load_workbook(SOURCE / "BEDE Murabaha Assessment .xlsx", read_only=True, data_only=True)
response = load_workbook(SOURCE / "Foo LMS Sharia Compliance.xlsx", read_only=True, data_only=True)

response_rows = {}
for row_number, values in enumerate(response["List Of Requirements"].iter_rows(values_only=True), 1):
    if isinstance(values[0], (int, float)):
        question = int(values[0])
        response_rows[question] = (row_number + 1, clean(response["List Of Requirements"].cell(row_number + 1, 2).value))


def theme(question):
    for last, label in [
        (5, "Product and pricing"),
        (11, "Lifecycle and ownership"),
        (17, "Schedule and contractual profit"),
        (21, "Ibra and early settlement"),
        (25, "Delinquency and rescheduling"),
        (34, "Accounting and provisioning"),
        (37, "Fees tax and takaful"),
        (40, "Documents and disclosure"),
        (44, "Governance audit and coexistence"),
        (47, "Scenario evidence"),
    ]:
        if question <= last:
            return label
    raise AssertionError(question)


def spec_reference(question):
    if question <= 5:
        return "Islamic Financing Module §§2–5; Pricing FRD §§3–6"
    if question <= 9:
        return "Lifecycle BRD §§5–10, 12, 18–19"
    if question == 10:
        return "Islamic Financing Module §12; Lifecycle BRD §19"
    if question == 11:
        return "Pricing FRD §§3, 6; Lifecycle BRD §19"
    if question <= 17:
        return "Pricing FRD §§5–6; Section 6 test workbook"
    if question <= 21:
        return "Ibra BRD BR-1–BR-7 and acceptance criteria"
    if question <= 25:
        return "Pricing FRD §6; Ibra BRD BR-5"
    if question <= 34:
        return "Islamic Financing Module §13; dedicated accounting/provisioning design not supplied"
    if question <= 37:
        return "Dedicated fee/tax/takaful design not supplied"
    if question <= 40:
        return "Islamic Financing Module §§4, 16; Lifecycle BRD §§9, 12"
    if question == 41:
        return "Formal Sharia approval/certification not supplied"
    if question <= 43:
        return "Lifecycle BRD §§8–10; Pricing FRD §6; Section 6 test workbook"
    if question == 44:
        return "Islamic Financing Module §§1–3"
    return "FOO response Scenario 1–3; Pricing FRD §6; Ibra BRD; Section 6 test workbook"


OPEN_CONTRADICTION = {6, 7, 18, 19}
STATUS_AMBIGUITY = {9, 14, 15, 16, 21, 22, 25, 40, 42, 43, 45, 46, 47}
FUTURE_OR_PARTIAL = {2, 8, 10, 23, 24, 26, 27, 28, 29, 30, 31, 32, 33, 34, 38, 39, 41}


def interpretation(question):
    if question in OPEN_CONTRADICTION:
        return "Open current-state contradiction: assessment says major gap; FOO answer claims support; BRD specifies enhancement"
    if question in STATUS_AMBIGUITY:
        return "Status/evidence ambiguity: reconcile present-tense or validated claim with assessment and repeatable proof"
    if question in FUTURE_OR_PARTIAL:
        return "Partial baseline or proposed extension: separate existing function from unbuilt requirement"
    return "Claimed baseline/configuration: verify on versioned build and compare with requirement"


def evidence_request(question):
    overrides = {
        6: "Versioned live lifecycle demo; persisted states and API contract; prove Phase 1 excludes MPO promise state or agree design change",
        7: "Negative UI/API tests before purchase, ownership, possession, evidence and acceptance; prove receivable stays inactive",
        9: "Persisted purchase/ownership/possession fields, required attachments, audit events and API payloads",
        18: "Versioned early-settlement demo; distinct Ibra transaction and ledger effect; conventional-loan regression",
        19: "Product configuration and each Ibra method; maker-checker pending/reject/approve evidence and calculation audit",
        21: "Sample settlement letter, statement and customer API showing Selling Price, Ibra and final settlement",
        41: "Documented Sharia board review and formal sign-off scope; do not infer certification from product design",
        45: "Rerun Scenario 1 with build/config ID, schedule and GL before/after month 12",
        46: "Rerun Scenario 2 with Ibra calculation, approvals, transaction and journal before/after",
        47: "Rerun Scenario 3 at 90 days past due with debt, accrual, charge and charity journal evidence",
    }
    if question in overrides:
        return overrides[question]
    if 12 <= question <= 17 or 20 <= question <= 25 or question == 43:
        return "Versioned product configuration plus executed positive, negative and conventional-regression tests; schedule and GL before/after"
    if 26 <= question <= 34:
        return "Approved accounting policy/posting matrix; GL mapping; journal samples, reconciliation and finance review"
    if 35 <= question <= 40:
        return "Approved business/Sharia rule, configured example and generated customer/report output"
    return "Versioned build/configuration, live demonstration or API result, persisted records and named reviewer sign-off"


matrix = []
for row_number, values in enumerate(assessment["Sheet1"].iter_rows(values_only=True), 1):
    if not isinstance(values[0], (int, float)):
        continue
    question = int(values[0])
    response_row, response_text = response_rows[question]
    matrix.append({
        "ID": f"Q{question:02d}",
        "Theme": theme(question),
        "Capability/question": clean(values[1]),
        "Assessment status": clean(values[3]),
        "Assessment baseline/FRD claim": clean(values[2]),
        "Assessment gap/comment": clean(values[4]),
        "FOO response claim": response_text,
        "Review interpretation": interpretation(question),
        "Related specification": spec_reference(question),
        "Evidence required": evidence_request(question),
        "Assessment reference": f"BEDE Murabaha Assessment / Sheet1!B{row_number}:E{row_number}",
        "FOO reference": f"Foo LMS Sharia Compliance / List Of Requirements!B{response_row}",
        "Disposition": "Open — current implementation unverified",
    })
assert len(matrix) == 47 and set(response_rows) == set(range(1, 48))
matrix_fields = list(matrix[0])
write_tsv(OUTPUT / "capability-matrix.tsv", matrix, matrix_fields)
(OUTPUT / "capability-matrix.json").write_text(json.dumps(matrix, indent=2, ensure_ascii=False) + "\n")

cases_book = load_workbook(SOURCE / "Murabaha_Section6_Test_Scenarios.xlsx", read_only=True, data_only=True)
cases = []
for row_number, values in enumerate(cases_book["Test Cases"].iter_rows(min_row=5, max_row=75, values_only=True), 5):
    if not clean(values[0]).startswith("TC-"):
        continue
    section = clean(values[1])
    section_number = int(section.split(".", 1)[1])
    wave = "1 — core profit invariant" if section_number in {3, 4, 5, 18} else (
        "2 — schedule and servicing" if 6 <= section_number <= 13 else "3 — controls and regression"
    )
    cases.append({
        "Test ID": clean(values[0]),
        "Wave": wave,
        "FRD section": section,
        "Area": clean(values[2]),
        "Scenario": clean(values[3]),
        "Type": clean(values[4]),
        "Baseline": clean(values[5]),
        "Preconditions": clean(values[6]),
        "Steps": clean(values[7]),
        "Test data": clean(values[8]),
        "Expected result": clean(values[9]),
        "FRD design-time status": clean(values[10]),
        "Priority in source": clean(values[11]),
        "Execution result": clean(values[13]),
        "Proposed executor": "FOO QA; Bede business/finance witness for high-risk results",
        "Evidence to retain": "Build/config ID; input and API/log evidence; schedule and GL before/after; screenshot/export; defect link",
        "Source reference": f"Murabaha_Section6_Test_Scenarios / Test Cases!A{row_number}:R{row_number}",
    })
assert len(cases) == 71 and all(case["Execution result"] == "Not Run" for case in cases)
write_tsv(OUTPUT / "section6-execution-tracker.tsv", cases, list(cases[0]))

targeted = [
    ("L01", "Lifecycle BRD §§5–6", "Separate Murabaha and financing states", "Approve financing, then progress required Murabaha stages", "Native financing remains approved; separate Murabaha state advances in order; no active receivable early"),
    ("L02", "Lifecycle BRD BR-2", "Block sale before ownership and possession", "Attempt sale via UI and API before purchase, ownership, possession, required evidence and acceptance", "Each premature attempt fails without sale or customer receivable"),
    ("L03", "Lifecycle BR-4–BR-5", "Mandatory evidence and audit", "Omit one required document, then provide it; inspect persisted event/history", "Sale blocked until all mandatory evidence exists; actors, timestamps and correction history retained"),
    ("L04", "Lifecycle BRD BR-3", "Maker-checker lifecycle transition", "Submit ownership confirmation, reject and resubmit/approve", "Pending/rejected actions do not advance state; approval advances once and audit retains both decisions"),
    ("L05", "Lifecycle BRD §§17–18", "Offer version and receivable activation", "Change commercial terms after offer, then try to use earlier acceptance and activate", "Earlier acceptance cannot execute changed terms; receivable activates only after valid sale execution"),
    ("L06", "Lifecycle BRD §§5, 19", "Phase 1 scope and MPO promise", "Inspect available lifecycle actions for direct Goods Murabaha and future MPO/Tawarruq options", "Phase 1 follows approved direct-Goods sequence; future promise/variant controls are disabled unless separately agreed"),
    ("L07", "Lifecycle BRD §20", "Conventional regression", "Run approval/disbursement on a conventional product", "Existing conventional lifecycle remains unchanged"),
    ("I01", "Ibra BRD BR-1", "Explicit Ibra transaction", "Settle an eligible Murabaha with a rebate and inspect transaction history", "Ibra is identified separately, reduces only remaining Profit and lowers settlement amount"),
    ("I02", "Ibra BRD BR-2", "Ibra calculation methods", "Calculate manual, percentage and unearned-profit methods with boundary values", "Each amount follows configured method and never exceeds eligible remaining Profit"),
    ("I03", "Ibra BRD BR-3", "Maker-checker and no early financial impact", "Submit, reject and approve an Ibra request; compare balances and audit", "Pending/rejected requests leave balances unchanged; approved request posts once"),
    ("I04", "Ibra BRD BR-3", "Customer self-service auto-approval", "Use system-calculated and overridden requests above/below configured cap", "Only eligible customer-channel, unmodified calculations within cap auto-approve"),
    ("I05", "Ibra BRD BR-6–BR-7", "Settlement quotation and API", "Generate back-office, customer API and document settlement views", "Remaining Selling Price, eligible/granted Ibra and final amount reconcile and use Islamic labels"),
    ("I06", "Ibra BRD BR-5 and acceptance criteria", "Partial settlement and arrears ordering", "Partially settle with overdue amounts and Recalculate Profit disabled", "Due/overdue amounts are paid first; Ibra applies only to permitted unearned Profit; remaining contract is unchanged"),
    ("I07", "Ibra BRD acceptance criteria", "Conventional regression", "Use prepay/interest waiver on a conventional loan", "Conventional terminology and behaviour remain unchanged"),
]
targeted_rows = [
    {"ID": id_, "Source section": source, "Verification target": target, "Proposed procedure": procedure,
     "Expected evidence/result": expected, "Owner": "Proposed: FOO QA; Bede product/finance reviewer",
     "Execution result": "Not Run", "Evidence link": "Pending execution", "Defect/decision": "None recorded"}
    for id_, source, target, procedure, expected in targeted
]
write_tsv(OUTPUT / "lifecycle-ibra-targeted-tests.tsv", targeted_rows, list(targeted_rows[0]))

summary = {
    "received_date": manifest["received_date"],
    "source_count": len(manifest["files"]),
    "mapped_questions": len(matrix),
    "assessment_status_counts": dict(Counter(item["Assessment status"] for item in matrix)),
    "open_direct_contradictions": [f"Q{question:02d}" for question in sorted(OPEN_CONTRADICTION)],
    "section6_cases": len(cases),
    "section6_execution_result_counts": dict(Counter(case["Execution result"] for case in cases)),
    "targeted_new_cases": len(targeted_rows),
    "note": "Source claims and proposed tests only; no LMS access, executed test, product sign-off or Sharia approval established.",
}
(OUTPUT / "review-metrics.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False) + "\n")
print(json.dumps({key: summary[key] for key in ("source_count", "mapped_questions", "section6_cases", "targeted_new_cases")}, ensure_ascii=False))
