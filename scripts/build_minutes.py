import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parent.parent
content=json.loads((root/'meetings/2026-10-01/supporting/minutes-content.json').read_text())
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.68)
sec.left_margin=sec.right_margin=Inches(.72)
styles=doc.styles
for name in ('Normal','Body Text','List Bullet','List Number'):
 s=styles[name]; s.font.name='Calibri'; s.font.size=Pt(10.5)
 s.paragraph_format.space_after=Pt(4)
 s.paragraph_format.line_spacing=1.03
for name,size in [('Title',23),('Heading 1',14),('Heading 2',11.5)]:
 s=styles[name]; s.font.name='Calibri'; s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
 s.paragraph_format.space_before=Pt(7); s.paragraph_format.space_after=Pt(4)
 s.paragraph_format.keep_with_next=True
styles['Title'].paragraph_format.space_before=Pt(0)
doc.add_paragraph(content['title'],'Title')
p=doc.add_paragraph(content['subtitle']); p.runs[0].font.color.rgb=RGBColor.from_string('555555')
for label,value in content['metadata']:
 p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(3); p.add_run(label+': ').bold=True; p.add_run(value)
doc.add_paragraph(content['overview'])
for section in content['sections']:
 if section.get('new_page'): doc.add_page_break()
 doc.add_heading(section['heading'],level=1)
 for entry in section.get('entries',[]):
  if isinstance(entry,str): doc.add_paragraph(entry)
  else:
   doc.add_heading(entry['title'],level=2); doc.add_paragraph(entry['text'])
 for bullet in section.get('bullets',[]): doc.add_paragraph(bullet,'List Bullet')
 if section.get('table'):
  table=doc.add_table(rows=1,cols=len(section['table']['headers'])); table.style='Table Grid'
  for i,h in enumerate(section['table']['headers']): table.rows[0].cells[i].text=h
  trpr=table.rows[0]._tr.get_or_add_trPr(); tag=OxmlElement('w:tblHeader'); trpr.append(tag)
  for row in section['table']['rows']:
   cells=table.add_row().cells
   for c,txt in zip(cells,row): c.text=txt
  for row in table.rows:
   trpr=row._tr.get_or_add_trPr(); tag=OxmlElement('w:cantSplit'); trpr.append(tag)
   for cell in row.cells:
    for p in cell.paragraphs:
     p.paragraph_format.space_after=Pt(2); p.paragraph_format.space_before=Pt(2)
     for r in p.runs:
      r.font.size=Pt(9.5); r.font.color.rgb=RGBColor(0,0,0); r.bold=False
  if section['table'].get('widths'):
   table.autofit=False
   for col,w in zip(table.columns,section['table']['widths']): col.width=Inches(w)
   for row in table.rows:
    for c,w in zip(row.cells,section['table']['widths']):c.width=Inches(w)
for table in doc.tables:
 for c in table.rows[0].cells:
  sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E9EDF2'); c._tc.get_or_add_tcPr().append(sh)
  for p in c.paragraphs:
   for r in p.runs:r.bold=True
for elem in (doc._element,doc.styles.element):
 for b in list(elem.iter(qn('w:pBdr'))): b.getparent().remove(b)
p=sec.footer.paragraphs[0]; p.alignment=2
r=p.add_run('Bede  |  Draft meetings/2026-10-01/minutes  |  '); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string('666666')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); p._p.append(fld)
doc.core_properties.title=content['title']; doc.core_properties.subject='Bede project planning meeting'
doc.core_properties.author=''; doc.core_properties.keywords='Bede, meeting meetings/2026-10-01/minutes, LMS, Mifos'
out=root/'meetings/2026-10-01/minutes/Bede_Meeting_Minutes_2026-10-01.docx'; doc.save(out)
md=['# '+content['title'],'',content['subtitle'],'']
md += [f'**{k}:** {v}  ' for k,v in content['metadata']]
md += ['',content['overview'],'']
for s in content['sections']:
 md += ['## '+s['heading'],'']
 for e in s.get('entries',[]):
  md += ([e,''] if isinstance(e,str) else ['### '+e['title'],'',e['text'],''])
 md += ['- '+b for b in s.get('bullets',[])]
 if s.get('table'):
  t=s['table']; md += ['', '| '+' | '.join(t['headers'])+' |','| '+' | '.join(['---']*len(t['headers']))+' |']
  md += ['| '+' | '.join(row)+' |' for row in t['rows']]
 md += ['']
(root/'meetings/2026-10-01/minutes/Bede_Meeting_Minutes_2026-10-01.md').write_text('\n'.join(md))
print(out)
