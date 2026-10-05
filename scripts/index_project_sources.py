from pathlib import Path
import hashlib
import json
import shutil
import argparse
from pypdf import PdfReader
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple
from pptx import Presentation

root = Path(__file__).resolve().parent.parent
parser = argparse.ArgumentParser(description='Preserve/index the supplied Bede documents without replacing different sources.')
parser.add_argument('--source-dir', type=Path, help='Optional directory containing original files; omit to index preserved project copies.')
args = parser.parse_args()
sources = [
    ('Annex2_LMS_Compliance_FOO_Response.xlsx', 'documentation/compliance', '2026-10-01'),
    ('FOO - Bede - Loan Management System - Technical Response V1.0.pdf', 'documentation/proposals', '2026-10-01'),
    ('FOO - Bede - Loan Management System - Commercial Response V1.6.pdf', 'documentation/proposals', '2026-10-01'),
    ('Bede LMS Kick-off.pdf', 'meetings/2026-10-01/client-kickoff', '2026-10-02'),
    ('Bede LMS Project Kickoff Deck.pptx', 'meetings/2026-10-01/client-kickoff', '2026-10-02'),
    ('Bede - Foo LMS - Scope and Requirements.pdf', 'documentation/requirements', '2026-10-02'),
    ('BEDE_Requirements_vs_FOO_Roadmap.xlsx', 'documentation/requirements', '2026-10-05'),
]
missing = [name for name, folder, _ in sources if not (root / folder / name).exists() and not (args.source_dir and (args.source_dir / name).exists())]
if missing:
    raise FileNotFoundError('Source files unavailable; provide durable copies before indexing: ' + ', '.join(missing))
extracted = root / 'documentation/search-index'
extracted.mkdir(parents=True, exist_ok=True)
previous_manifest_path = extracted / 'source-manifest.json'
previous_entries = {entry['name']: entry for entry in json.loads(previous_manifest_path.read_text()).get('files', [])} if previous_manifest_path.exists() else {}
manifest = {'timezone': 'Asia/Beirut', 'files': []}
for name, folder, received in sources:
    destination = root / folder
    destination.mkdir(parents=True, exist_ok=True)
    saved = destination / name
    candidate = args.source_dir / name if args.source_dir else saved
    original = candidate if candidate.exists() else saved
    original_hash = hashlib.sha256(original.read_bytes()).hexdigest()
    if saved.exists():
        if hashlib.sha256(saved.read_bytes()).hexdigest() != original_hash:
            raise RuntimeError(f'Existing source differs: {saved}')
    else:
        shutil.copy2(original, saved)
    assert hashlib.sha256(saved.read_bytes()).hexdigest() == original_hash
    original_path = previous_entries.get(name, {}).get('original_path', str(candidate) if candidate.exists() and args.source_dir else 'Preserved project copy')
    entry = {'name': name, 'received_date': received, 'original_path': original_path, 'project_path': saved.relative_to(root).as_posix(), 'sha256': original_hash, 'bytes': saved.stat().st_size}
    if candidate.exists() and args.source_dir and str(candidate) != original_path:
        entry['available_source_path'] = str(candidate)
    elif previous_entries.get(name, {}).get('available_source_path'):
        entry['available_source_path'] = previous_entries[name]['available_source_path']
    if saved.suffix == '.pdf':
        pdf = PdfReader(saved)
        text_path = extracted / f'{saved.stem}.txt'
        text_path.write_text('\n\n'.join(f'=== PDF PAGE {i} ===\n{page.extract_text() or ""}' for i, page in enumerate(pdf.pages, 1)))
        entry.update(pages=len(pdf.pages), searchable_text=text_path.relative_to(root).as_posix())
        print(f'{name}: {len(pdf.pages)} pages')
    elif saved.suffix == '.pptx':
        deck = Presentation(saved)
        text_path = extracted / f'{saved.stem}.txt'
        slide_text = []
        for i, slide in enumerate(deck.slides, 1):
            slide_text.append(f'=== SLIDE {i} ===\n' + '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame and shape.text.strip()))
        text_path.write_text('\n\n'.join(slide_text))
        entry.update(slides=len(deck.slides), searchable_text=text_path.relative_to(root).as_posix())
        print(f'{name}: {len(deck.slides)} slides')
    else:
        book = load_workbook(saved, read_only=True, data_only=False)
        sheets = []
        output = {}
        for sheet in book:
            rows = []
            for row in sheet.iter_rows():
                cells = {cell.coordinate: cell.value for cell in row if cell.value is not None}
                if cells:
                    rows.append(cells)
            output[sheet.title] = rows
            bounds = [coordinate_to_tuple(cell) for row in rows for cell in row]
            sheets.append({'title': sheet.title, 'state': sheet.sheet_state, 'populated_last_row': max((r for r, c in bounds), default=0), 'populated_last_column': max((c for r, c in bounds), default=0), 'nonempty_rows': len(rows)})
            print(f'{sheet.title}: rows={sheet.max_row}, cols={sheet.max_column}, populated rows={len(rows)}')
            print(json.dumps(rows[:6], ensure_ascii=False, default=str))
        book.close()
        text_path = extracted / f'{saved.stem}.json'
        text_path.write_text(json.dumps(output, indent=2, ensure_ascii=False, default=str))
        entry.update(sheets=sheets, searchable_cells=text_path.relative_to(root).as_posix())
    manifest['files'].append(entry)
(extracted / 'source-manifest.json').write_text(json.dumps(manifest, indent=2))
